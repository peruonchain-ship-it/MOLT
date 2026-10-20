import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from .config import TIMEZONE
from .matching import FAMILIES

CLASS_LABELS = {"open": "Vigente según fuente", "review": "Por verificar", "historical": "Histórico"}
HEADERS = ["Entidad", "Objeto", "Referencia", "Zona", "Base de ubicación", "Estado", "Motivo del estado",
           "Monto referencial", "Moneda", "Fecha límite (origen)", "Publicación", "Familias detectadas",
           "Alcance comercial", "Coincidencias y evidencia", "Revisión de documentos", "Etapa comercial", "Notas",
           "Adjudicatarios (sin asumir subcontratación)", "Enlace oficial", "Datos OCDS", "Consultado por MOLT"]


def spreadsheet(items, source_status=None):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Oportunidades"
    sheet.append(HEADERS)
    for item in items:
        docs = item.get("document_checks") or []
        row = [item["buyer"], item["title"], item["reference"], item["city"] or "Sin confirmar",
               item["location_basis"], CLASS_LABELS[item["classification"]], item["status_reason"], item["amount"],
               item["currency"], item["deadline"] or "No publicada", item["published"] or "No publicada",
               ", ".join(FAMILIES[f][0] for f in item["families"]), item["scope"],
               " | ".join(f"{e['origin']}: {e['term']} — {e['snippet']}" for e in item["evidence"]),
               " | ".join(f"{d['title']}: {d['status']} — {d['detail']}" for d in docs) or "Pendiente/no disponible",
               item.get("stage", "Pendiente"), item.get("notes", ""), "; ".join(item.get("suppliers", [])),
               item["url"], item["raw_url"], item.get("last_seen", "")]
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", str(v))[:32000] if isinstance(v, str) else v for v in row])
        # All imported strings are text, even if they begin with =, +, - or @.
        for cell, value in zip(sheet[sheet.max_row], row, strict=True):
            if isinstance(value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="123B32")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    widths = [38, 85, 32, 18, 50, 24, 65, 20, 12, 30, 30, 44, 45, 90, 65, 20, 55, 45, 60, 60, 30]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[7].number_format = '#,##0.00'
    meta = wb.create_sheet("Fuente y alcance")
    meta.append(["MOLT · radar privado", "Exportación de registros, no garantía de vigencia ni adjudicación."])
    meta.append(["Generado (Perú)", datetime.now(ZoneInfo(TIMEZONE)).isoformat()])
    meta.append(["Fuente", "https://contratacionesabiertas.oece.gob.pe/api"])
    if source_status:
        for key in ("status", "window_start", "window_end", "scanned", "source_latest", "message"):
            meta.append([key, str(source_status.get(key, ""))])
    meta.append(["Ubicación", "El domicilio de la entidad es referencial; confirmar lugar de prestación."])
    meta.append(["Documentos", "Solo PDFs con texto, hasta 160 páginas/400.000 caracteres. Escaneados y ZIP: revisión manual."])
    meta.append(["Contratos integrales", "Una coincidencia indica un componente posible; revisar equipos, cantidades y requisitos."])
    meta.column_dimensions["A"].width = 30
    meta.column_dimensions["B"].width = 110
    for row in meta:
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)[:32000]
                cell.data_type = "s"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
