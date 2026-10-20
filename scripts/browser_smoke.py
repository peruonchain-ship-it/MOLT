"""Browser smoke test in a disposable instance. Does not change the real owner's account or send mail."""

import io
import json
import os
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import httpx
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parent.parent


def main():
    with tempfile.TemporaryDirectory(prefix="molt-browser-") as temp:
        folder = Path(temp)
        # Snapshot public records from the validated instance, but never copy owner credentials or sessions.
        from molt.config import Config
        from molt.db import Database

        db = Database(Config(data_dir=folder, worker_enabled=False, mail_to="browser-test@example.com"))
        source = ROOT / ".data/molt.sqlite3"
        if not source.exists():
            raise SystemExit("Ejecuta primero una consulta real: uv run python -m molt sync")
        with sqlite3.connect(source) as origin, db.connect() as target:
            for table in ("opportunities", "documents", "sync_runs"):
                for row in origin.execute(f"SELECT * FROM {table}"):
                    target.execute(f"INSERT INTO {table} VALUES ({','.join('?' for _ in row)})", row)
            target.execute("DELETE FROM sync_runs WHERE status='running'")
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        env = {**os.environ, "MOLT_DATA_DIR": temp, "MOLT_WORKER_ENABLED": "false", "MOLT_PORT": str(port),
               "MOLT_HOST": "127.0.0.1", "RESEND_API_KEY": "", "MOLT_MAIL_FROM": ""}
        log_file = (folder / "server.log").open("w")
        server = subprocess.Popen([sys.executable, "-m", "molt", "serve"], cwd=ROOT, env=env, stdout=log_file, stderr=log_file)
        base = f"http://127.0.0.1:{port}"
        try:
            for _ in range(100):
                try:
                    if httpx.get(base + "/health", timeout=1).status_code == 200:
                        break
                except httpx.RequestError:
                    pass
                time.sleep(0.1)
            else:
                raise RuntimeError("El servidor de prueba no arrancó.")
            with sync_playwright() as p:
                browser = p.chromium.launch(executable_path="/usr/bin/chromium", args=["--no-sandbox"])
                page = browser.new_page(viewport={"width": 1512, "height": 1100}, device_scale_factor=1)
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto(base + "/app")
                page.locator("#auth-screen").wait_for(state="visible")
                page.locator("#setup-code").fill((folder / "setup-code").read_text())
                page.locator("#username").fill("browser-owner")
                page.locator("#password").fill("browser-test-password-123")
                page.locator("#auth-submit").click()
                page.locator("#application").wait_for(state="visible")
                page.locator(".opportunity").first.wait_for()
                total = int(page.locator("#result-count").inner_text().replace(",", ""))
                assert total > 0
                page.screenshot(path="/tmp/molt-desktop.png", full_page=False)
                # Search, real detail, and a persisted commercial note.
                page.locator("#search").fill("zz-no-match-browser-test")
                page.locator(".empty-state").wait_for()
                page.locator("#search").fill("")
                page.locator(".opportunity").first.wait_for()
                page.locator(".card-action").first.click()
                page.locator("#detail-dialog").wait_for(state="visible")
                page.locator("#detail-stage").select_option("Cotizar")
                page.locator("#detail-notes").fill("Nota privada de prueba en base temporal")
                page.get_by_role("button", name="Guardar seguimiento").click()
                page.locator("#toast").filter(has_text="Seguimiento guardado").wait_for()
                page.locator("#close-detail").click()
                page.locator('.nav-item[data-view="followup"]').click()
                expect(page.locator("#result-count")).to_have_text("1")
                with page.expect_download() as download_info:
                    page.locator("#export-button").click()
                download = download_info.value
                workbook = load_workbook(io.BytesIO(Path(download.path()).read_bytes()))
                assert workbook.active.max_row == 2
                assert workbook.active.cell(2, 17).value == "Nota privada de prueba en base temporal"
                page.locator('.nav-item[data-view="settings"]').click()
                page.locator("#email-to").fill("private-test@example.com")
                page.get_by_role("button", name="Guardar preferencias").click()
                page.locator("#settings-message").filter(has_text="Preferencias guardadas").wait_for()
                page.locator('.nav-item[data-view="mail"]').click()
                page.get_by_role("button", name="Ver próximo resumen").click()
                page.locator("#mail-preview").wait_for(state="visible")
                assert page.locator("#send-mail").is_disabled()
                page.locator('.nav-item[data-view="radar"]').click()
                page.set_viewport_size({"width": 390, "height": 844})
                page.screenshot(path="/tmp/molt-mobile.png", full_page=False)
                overflow = page.evaluate("""() => [...document.querySelectorAll('body *')]
                    .filter(el => !el.closest('[hidden]') && el.getBoundingClientRect().right > innerWidth + 1)
                    .map(el => ({tag:el.tagName, id:el.id, cls:el.className,
                        right:el.getBoundingClientRect().right})).slice(0,10)""")
                assert page.evaluate("() => document.documentElement.scrollWidth <= innerWidth + 1"), overflow
                page.locator(".card-action").first.click()
                page.locator("#detail-dialog").wait_for(state="visible")
                page.locator("#close-detail").click()
                page.set_viewport_size({"width": 1512, "height": 1100})
                page.locator("#logout").click()
                page.locator("#auth-screen").wait_for(state="visible")
                assert page.request.get(base + "/api/dashboard").status == 401
                assert errors == [], errors
                browser.close()
            print(json.dumps({"browser": "passed", "real_records_visible": total,
                              "checks": ["private setup", "search", "detail", "follow-up", "filtered Excel",
                                         "preferences", "mail preview", "mobile layout", "logout"],
                              "external_emails_sent": 0}, ensure_ascii=False))
        finally:
            server.terminate()
            try:
                server.wait(timeout=10)
            except subprocess.TimeoutExpired:
                server.kill()
            log_file.close()


if __name__ == "__main__":
    main()
