import io

from openpyxl import load_workbook

from molt.source import save_release


def test_private_data_requires_session(client):
    for path in ["/api/dashboard", "/api/opportunities", "/api/export.xlsx", "/api/settings", "/api/mail/preview"]:
        assert client.get(path).status_code == 401
    assert client.get("/health").json()["status"] == "ok"
    assert client.get("/").status_code == 200


def test_setup_code_is_required_and_removed_after_use(client, config):
    bad = client.post("/api/auth/setup", json={"username": "owner", "password": "private-password-123", "code": "incorrect-code-123"})
    assert bad.status_code == 403
    code = (config.data_dir / "setup-code").read_text()
    good = client.post("/api/auth/setup", json={"username": "owner", "password": "private-password-123", "code": code})
    assert good.status_code == 200
    assert not (config.data_dir / "setup-code").exists()
    assert "HttpOnly" in good.headers["set-cookie"] and "SameSite=strict" in good.headers["set-cookie"]
    assert client.post("/api/auth/setup", json={"username": "second", "password": "private-password-123", "code": code}).status_code == 409


def test_csrf_followup_and_export_use_same_filter(owner, release):
    save_release(owner.app.state.db, release)
    items = owner.get("/api/opportunities?city=Lima&classification=open").json()
    assert items["total"] == 1
    response = owner.patch("/api/opportunities/" + release["ocid"], json={"stage": "Cotizar", "notes": "Revisar pantalla"})
    assert response.status_code == 200
    assert owner.get("/api/opportunities?followup=true").json()["total"] == 1
    token = owner.headers.pop("X-CSRF-Token")
    assert owner.patch("/api/opportunities/" + release["ocid"], json={"stage": "Descartada"}).status_code == 403
    owner.headers["X-CSRF-Token"] = token
    response = owner.get("/api/export.xlsx?city=Arequipa")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 1
    response = owner.get("/api/export.xlsx?city=Lima")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 2 and sheet.cell(2, 16).value == "Cotizar"


def test_invalid_filters_and_settings_are_rejected(owner):
    assert owner.get("/api/opportunities?city=Invalid").status_code == 422
    assert owner.get("/api/opportunities?family=not-a-family").status_code == 422
    settings = owner.get("/api/settings").json()["settings"]
    assert owner.put("/api/settings", json={**settings, "email_enabled": True}).status_code == 400
    assert owner.put("/api/settings", json={**settings, "cities": []}).status_code == 422
    assert owner.put("/api/settings", json={**settings, "email_to": "bad\r\nBcc:attacker@example.com"}).status_code == 422
    assert owner.put("/api/settings", json={**settings, "cities": ["Lima", "Trujillo"]}).status_code == 200


def test_logout_invalidates_session(owner):
    assert owner.post("/api/auth/logout").status_code == 200
    assert owner.get("/api/dashboard").status_code == 401


def test_password_change_invalidates_old_session(owner):
    old_cookie = owner.cookies.get("molt_session")
    response = owner.post("/api/auth/password", json={"current_password": "private-test-password-123", "new_password": "a-new-private-password-123"})
    assert response.status_code == 200
    owner.cookies.set("molt_session", old_cookie)
    assert owner.get("/api/auth/me").status_code == 401


def test_excel_strings_cannot_become_formulas(owner, release):
    release["buyer"]["name"] = '=HYPERLINK("https://invalid.example","click")'
    save_release(owner.app.state.db, release)
    content = owner.get("/api/export.xlsx").content
    wb = load_workbook(io.BytesIO(content), data_only=False)
    assert wb.active.cell(2, 1).data_type == "s"
    assert wb.active.cell(2, 1).value.startswith("=")
    assert all(cell.data_type != "f" for row in wb.active for cell in row)
