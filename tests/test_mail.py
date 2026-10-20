import base64
import io
import json
from dataclasses import replace
from datetime import UTC, datetime, timedelta

import httpx
import pytest
from openpyxl import load_workbook

from molt.db import Database, utcnow
from molt.mail import build_payload, digest_items, send_digest
from molt.source import save_release


@pytest.fixture
def mail_db(config, release):
    config = replace(config, mail_key="test-key-not-real", mail_from="radar@example.com")
    db = Database(config)
    save_release(db, release)
    with db.connect() as conn:
        conn.execute("INSERT INTO sync_runs(started_at,finished_at,status) VALUES (?,?,'complete')", (utcnow(), utcnow()))
    return db


def test_payload_contains_real_excel_and_fixed_recipient(mail_db):
    items = digest_items(mail_db)
    payload = build_payload(mail_db, items)
    assert payload["to"] == ["owner@example.com"]
    workbook = load_workbook(io.BytesIO(base64.b64decode(payload["attachments"][0]["content"])))
    assert workbook.active.max_row == 2
    assert "test-key-not-real" not in json.dumps(payload)


def test_delivery_is_recorded_only_after_provider_success(mail_db):
    calls = []

    def provider(request):
        calls.append(request)
        return httpx.Response(200, json={"id": "provider-test-1"})

    transport = httpx.MockTransport(provider)
    result = send_digest(mail_db, transport=transport)
    assert result["status"] == "sent" and result["count"] == 1
    assert len(calls) == 1 and calls[0].headers["Idempotency-Key"].startswith("molt-")
    assert send_digest(mail_db, transport=transport)["status"] == "empty"
    assert len(calls) == 1


def test_provider_failure_preserves_pending_items_and_retry_payload(mail_db):
    requests = []

    def provider(request):
        requests.append(request)
        return httpx.Response(503) if len(requests) == 1 else httpx.Response(200, json={"id": "accepted"})

    transport = httpx.MockTransport(provider)
    assert send_digest(mail_db, transport=transport)["status"] == "error"
    assert len(digest_items(mail_db)) == 1
    assert send_digest(mail_db, transport=transport)["status"] == "sent"
    assert requests[0].headers["Idempotency-Key"] == requests[1].headers["Idempotency-Key"]
    assert requests[0].content == requests[1].content


def test_historical_not_notified_by_default(mail_db, release):
    release["tender"]["status"] = "complete"
    release["contracts"] = [{"id": "closed"}]
    save_release(mail_db, release)
    assert digest_items(mail_db) == []
    mail_db.save_settings({"include_historical": True})
    assert len(digest_items(mail_db)) == 1


def test_imminent_deadline_reminder_is_once_per_day(mail_db, release):
    release["tender"]["tenderPeriod"]["endDate"] = (datetime.now(UTC) + timedelta(hours=24)).isoformat()
    save_release(mail_db, release)
    transport = httpx.MockTransport(lambda _: httpx.Response(200, json={"id": "accepted"}))
    assert send_digest(mail_db, transport=transport)["status"] == "sent"
    assert digest_items(mail_db) == []
    with mail_db.connect() as conn:
        conn.execute("UPDATE deliveries SET sent_at=?", ((datetime.now(UTC) - timedelta(days=1)).isoformat(),))
    reminder = digest_items(mail_db)
    assert len(reminder) == 1 and ":reminder:" in reminder[0]["delivery_key"]
    assert send_digest(mail_db, transport=transport)["status"] == "sent"
    assert digest_items(mail_db) == []


def test_missing_credentials_or_incomplete_source_blocks_send(mail_db):
    with mail_db.connect() as conn:
        conn.execute("UPDATE sync_runs SET status='partial'")
    result = send_digest(mail_db)
    assert result["status"] == "blocked"
    assert "consulta completa" in result["message"]


def test_missing_key_does_not_attempt_external_send(db):
    assert send_digest(db)["status"] == "blocked"
